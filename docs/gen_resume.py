# -*- coding: utf-8 -*-
"""智构 StructMind · 人机协同履历表生成（按官方模板结构）
以官方模板为底，填入真实信息：
  姓名=刘宇翔 / AI使用时间=3年 / 勾选 L3-L6 / 5 个真实使用场景（含编程智能体）
输出：人机协同履历表_智构StructMind.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = r"D:\电脑管家迁移文件\微信聊天记录搬家\xwechat_files\wxid_s01sqpdvf67s22_b28e\msg\file\2026-09\人机协同履历表模板.docx"
OUT = os.path.join(BASE, "人机协同履历表_智构StructMind.docx")

DEEP = RGBColor(0x0F, 0x4C, 0x81)
INK = RGBColor(0x20, 0x2A, 0x33)


def set_run(run, text, size=10.5, bold=False, color=INK, east="宋体"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        from docx.oxml import OxmlElement
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), east)


def fill_cell(cell, lines, size=10.5, bold=False, color=INK, east="宋体"):
    """清空单元格后用给定行列表重填（每行一个段落）"""
    cell.text = ""
    p0 = cell.paragraphs[0]
    for i, line in enumerate(lines):
        p = p0 if i == 0 else cell.add_paragraph()
        r = p.add_run()
        set_run(r, line, size=size, bold=bold, color=color, east=east)


CAPABILITY = [
    "请根据L1至L6六个能力层级进行自我评估（勾选）：",
    "□ L1基础问答：能查询信息、解释概念、生成简单文本",
    "□ L2提示词：能设置角色、目标、格式，提升输出质量",
    "☑ L3工作辅助：能将AI融入分析、汇报与创意等任务",
    "☑ L4流程自动化：能设计固定流程减少重复劳动",
    "☑ L5智能体：能搭建面向具体任务的智能体",
    "☑ L6业务创新：能基于AI重新设计业务流程或交付方式",
    "自评说明：3 年间从「会用 AI 问答」起步，到能设计多 Agent 协同架构、用编程智能体完成全流程开发与自动化验证，最终交付「智构 StructMind」参赛作品。",
]

SCENARIOS = [
    ("智能体搭建（核心）", "WorkBuddy（编程智能体）",
     "与编程智能体多轮协作完成「智构 StructMind」全流程：需求拆解→多 Agent 架构设计（DAG 编排引擎/主动记忆系统/元认知反思）→全部代码实现→9 套 312 项自动化回归验证→GitHub Pages 部署上线。我负责架构决策、工程语义把关与最终验证合入。"),
    ("代码开发与调试", "WorkBuddy（编程智能体）",
     "多 Agent 管线（pipeline/real-engine/optimizer/dag-engine/memory/metacognition）、规范校核规则引擎、3D 结构可视化、UI 设计系统等代码由编程智能体生成，我逐项复核并人工验证工程正确性。累计 95 次 git 提交全程可追溯。"),
    ("技术评审与优化", "DeepSeek",
     "以「评委视角」提供 4 类 20+ 项优化意见（底层逻辑/功能完整性/视觉包装/答辩防翻车），我逐条甄别（取其精华去其糟粕）后分 3 批落地并回归验证。"),
    ("知识库整理", "豆包 / DeepSeek",
     "梳理建筑规范条文（GB 55002/GB/T 50011 等）、碳排放因子（GB/T 51366）等结构化知识库材料，我按规范原文逐条校验工程语义。"),
    ("文档与答辩材料", "豆包 / WorkBuddy",
     "技术说明文档、演示视频脚本、人机协同履历表、答辩 Q&A 预案的起草与润色。"),
]


def main():
    doc = Document(TEMPLATE)
    t = doc.tables[0]

    # --- 行0：姓名 / AI使用时间 ---
    fill_cell(t.rows[0].cells[1], ["刘宇翔"], size=11, bold=True, color=DEEP)
    fill_cell(t.rows[0].cells[4], ["3年"], size=11, bold=True, color=DEEP)

    # --- 行1：使用能力评估（合并单元格，5 列同源）---
    seen = set()
    for c in t.rows[1].cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        fill_cell(c, CAPABILITY, size=10)

    # --- 行2：表头 使用场景/AI工具/描述 ---
    fill_cell(t.rows[2].cells[0], ["使用场景"], size=10.5, bold=True, color=DEEP)
    fill_cell(t.rows[2].cells[1], ["AI工具"], size=10.5, bold=True, color=DEEP)
    fill_cell(t.rows[2].cells[2], ["描述"], size=10.5, bold=True, color=DEEP)

    # --- 行3~7：场景数据（共 5 条，模板有 3+2 行可用）---
    data_rows = [t.rows[3], t.rows[4], t.rows[5], t.rows[6], t.rows[7]]
    for idx, row in enumerate(data_rows):
        if idx < len(SCENARIOS):
            name, tool, desc = SCENARIOS[idx]
            # 描述列跨 3 列合并
            seen = set()
            for j, c in enumerate(row.cells):
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                if j == 0:
                    fill_cell(c, [name], size=10)
                elif j == 1:
                    fill_cell(c, [tool], size=10, color=DEEP, bold=True)
                else:
                    fill_cell(c, [desc], size=9.5)
        else:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                fill_cell(c, [""], size=10)

    # --- 行8：补充说明（项目分工与验证过程）---
    seen = set()
    for c in t.rows[8].cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        fill_cell(c, [
            "项目分工与验证过程",
            "本人负责：创意与需求定义、多智能体架构决策（规划/执行分离、辩论回退、强条一票否决）、结构专业知识把关、规范条文工程语义复核、每次 AI 产出的验证合入决策、现场演示答辩。",
            "AI 负责：系统架构设计辅助、代码实现、视觉设计、自动化测试、部署发布。",
            "修改验证闭环：AI 改 → 全量回归（typecheck / eslint / 9 套 312 项断言）→ 线上核验 → 人工复核。涉及结构计算的修改一律由本人按现行规范人工复算确认后才上线；涉及线上发布的修改均以实际访问线上产物核对关键功能为准。历次修改全程 git 可追溯。",
        ], size=9.5)

    doc.save(OUT)
    print("SAVED ->", OUT)


if __name__ == "__main__":
    main()
